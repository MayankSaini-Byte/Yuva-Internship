"""Regenerate the Word report from outputs produced by preprocess.py."""
from pathlib import Path
import json
import pandas as pd
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parent
OUT = ROOT / 'outputs'
a = json.loads((OUT / 'audit_summary.json').read_text())
stats = pd.read_csv(OUT / 'initial_statistics.csv', index_col=0)
doc = Document()
sec = doc.sections[0]
sec.page_width, sec.page_height = Inches(8.5), Inches(11)
sec.top_margin = sec.bottom_margin = Inches(.7)
sec.left_margin = sec.right_margin = Inches(.8)
for name in ['Normal', 'Title', 'Subtitle', 'Heading 1', 'Heading 2']:
    style = doc.styles[name]
    style.font.name = 'Calibri'
    style.font.color.rgb = RGBColor(0, 0, 0)
doc.styles['Normal'].font.size = Pt(10.5)
doc.styles['Normal'].paragraph_format.space_after = Pt(7)
doc.styles['Normal'].paragraph_format.line_spacing = 1.08
doc.styles['Title'].font.size = Pt(26)
doc.styles['Heading 1'].font.size = Pt(18)
doc.styles['Heading 2'].font.size = Pt(12)
header = sec.header.paragraphs[0]
header.text = 'WEEK 1    |    AUTO MPG DATA PREPARATION'
header.runs[0].font.size = Pt(8)
footer = sec.footer.paragraphs[0]
footer.alignment = 2
footer.add_run('Week 1 report  |  ')
field = OxmlElement('w:fldSimple'); field.set(qn('w:instr'), 'PAGE'); footer._p.append(field)

def p(text): doc.add_paragraph(text)
def h(text): doc.add_heading(text, 2)
def page(title): doc.add_page_break(); doc.add_heading(title, 1)
def code(text):
    para = doc.add_paragraph()
    para.paragraph_format.space_after = Pt(9)
    para.paragraph_format.line_spacing = 1
    for i, line in enumerate(text.strip().splitlines()):
        run = para.add_run(('\n' if i else '') + line)
        run.font.name = 'DejaVu Sans Mono'; run.font.size = Pt(8)

def table(headers, rows, widths=None):
    t = doc.add_table(rows=1, cols=len(headers))
    t.autofit = False
    if widths:
        for col, width in zip(t.columns, widths): col.width = Inches(width)
    for c, text in zip(t.rows[0].cells, headers): c.text = str(text)
    repeat = OxmlElement('w:tblHeader'); t.rows[0]._tr.get_or_add_trPr().append(repeat)
    for row in rows:
        for c, text in zip(t.add_row().cells, row): c.text = str(text)
    borders = OxmlElement('w:tblBorders')
    for side in ['top', 'left', 'bottom', 'right', 'insideH', 'insideV']:
        el=OxmlElement('w:'+side); el.set(qn('w:val'),'single'); el.set(qn('w:sz'),'4'); el.set(qn('w:color'),'D9D9D9'); borders.append(el)
    t._tbl.tblPr.append(borders)
    for i,row in enumerate(t.rows):
        for j,c in enumerate(row.cells):
            if widths: c.width=Inches(widths[j])
            props=c._tc.get_or_add_tcPr()
            shade=OxmlElement('w:shd'); shade.set(qn('w:fill'), 'DCE6EC' if i==0 else 'FFFFFF'); props.append(shade)
            margins=OxmlElement('w:tcMar')
            for side in ['top','left','bottom','right']:
                el=OxmlElement('w:'+side); el.set(qn('w:w'),'75'); el.set(qn('w:type'),'dxa'); margins.append(el)
            props.append(margins)
            for para in c.paragraphs:
                para.paragraph_format.space_after=Pt(3)
                para.paragraph_format.space_before=Pt(3)
                for run in para.runs: run.font.size=Pt(9); run.bold=i==0
    doc.add_paragraph().paragraph_format.space_after=Pt(1)

doc.add_heading('Auto MPG Data Acquisition Cleaning and Preprocessing', 0)
p('Week 1 Project Report')
p('Prepared for academic submission')
h('Project objective')
p('This project prepares a public vehicle fuel-efficiency dataset for further analysis using Python. The workflow covers acquisition, structural inspection, missing values, inconsistent text, numerical validation, outlier review and reproducible feature preparation. The submitted program was executed against the included source data, and the tables and figures in this report describe that run.')
h('Main results')
table(['Measure', 'Verified result'], [
('Source records and fields', f'{a["rows_raw"]} records and {a["columns_raw"]} fields'),
('Missing horsepower', '6 entries filled using the training median of 92'),
('Vehicle names standardized', f'{a["names_changed"]} records with a retained correction log'),
('Exact duplicate records', str(a['duplicates'])),
('Numerical domain violations', '0 under the documented validation rules'),
('Potential outlier records', f'{a["outlier_rows"]} flagged and retained'),
('Prepared partitions', '318 training rows and 80 test rows'),
('Prepared predictors', '14 numerical features with no missing values')], [3.1,3.7])
h('Scope and outcome')
p('All 398 records were retained. Missing values were treated without using test-set statistics, while suspicious but plausible values were preserved for later review. The output includes readable cleaned data, separate feature and target files, a reusable preprocessing object and an audit trail. No predictive model is trained in Week 1, so no improvement in prediction accuracy is claimed.')
p('The assignment estimates 30 to 35 hours for the overall activity. This report does not represent that estimate as a measured work log.')

page('1 Dataset selection and acquisition')
p('I selected Auto MPG because its small size makes every cleaning decision inspectable, while its missing measurements and naming inconsistencies provide practical cleaning tasks. The dataset concerns city-cycle fuel consumption. UCI documents 398 records, seven predictors, a target and a vehicle-name identifier. Its published version already excludes eight records with unknown MPG. [1]')
p('A Kaggle listing is available at the link in Reference 2. The submitted data were downloaded directly from UCI, avoiding a Kaggle account dependency. The original text file and its supplied documentation are included. No synthetic errors or extra records were introduced.')
h('Collection method')
p('The source uses whitespace-separated numeric fields and quoted vehicle names. Parsing respects quotation marks so multiword names remain in one field. The program first imports strings, preserving the question-mark missing-value marker for auditing, then converts the numerical fields.')
code('''raw = pd.read_csv(raw_path, sep=r'\\s+', names=COLS,
                  quotechar='"', dtype=str)
if raw.shape != (398, 9):
    raise ValueError(f'Unexpected source shape: {raw.shape}')
df = raw.replace('?', np.nan).copy()''')
h('Field dictionary')
table(['Field', 'Meaning and treatment'], [
('mpg', 'Fuel efficiency in miles per gallon; target, kept unscaled'),
('cylinders', 'Cylinder-count category; one-hot encoded'),
('displacement', 'Engine displacement; numeric, source scale retained'),
('horsepower', 'Engine power; missing values imputed, then scaled'),
('weight', 'Vehicle weight; numeric, source scale retained'),
('acceleration', 'Acceleration measure; numeric, source scale retained'),
('model_year', 'Two-digit year code 70 to 82; treated as ordered numeric'),
('origin', 'Origin category coded 1 to 3; one-hot encoded'),
('car_name', 'Vehicle label; standardized and excluded from model inputs')], [1.5,5.3])
p('Units not specified in the supplied field documentation are not converted or inferred. The dataset is credited to R. Quinlan through UCI; the repository lists a CC BY 4.0 license. [1]')

page('2 Initial exploration and quality audit')
p('The initial audit inspects dimensions, missing counts, numeric conversion, duplicate records, value ranges and distributions. A stable row_id records the original one-based row position and links outputs back to the source; it is never used as a predictor.')
code('''audit['missing_before'] = df.isna().sum().to_dict()
audit['duplicates'] = int(df.duplicated().sum())
df[COLS[:-1]].describe().round(4).to_csv(
    out / 'initial_statistics.csv')''')
h('Observed numerical summary')
rows=[]
for c in ['mpg','displacement','horsepower','weight','acceleration']:
    rows.append([c, int(stats.loc['count',c]), f"{stats.loc['min',c]:g}", f"{stats.loc['mean',c]:.2f}", f"{stats.loc['max',c]:g}"])
table(['Field','Present','Minimum','Mean','Maximum'],rows,[1.6,.8,1.1,1.1,1.2])
p('Horsepower is present for 392 of 398 vehicles. Its six missing entries represent 1.51% of records. All other source fields are complete after parsing, and there are no exact duplicate rows. Numeric values passed the chosen domain checks; this is evidence of consistency with those checks, not proof that every measurement is correct.')
doc.add_picture(str(OUT/'exploration.png'),width=Inches(6.8))
p('Figure 1. Observed horsepower and the weight–MPG relationship, generated from the included data. The dashed line is the training-set median used for imputation. The scatter plot shows an association; it does not establish causation.')
p('Horsepower has a longer upper tail, which motivates a median-based fill. The plots summarize the complete dataset for documentation only. No data-dependent imputation, scaling or outlier threshold is fitted on the combined partitions.')

page('3 Missing values and inconsistent entries')
h('Missing horsepower')
p('Dropping every incomplete record would discard six vehicles with otherwise usable information. Instead, the data are split with test_size=0.20 and random_state=42. The training partition contains five missing horsepower values and the test partition contains one. The median of observed training horsepower is 92, which fills all six entries. A horsepower_missing flag preserves which observations were imputed.')
code('''train_ids, test_ids = train_test_split(
    df.index, test_size=0.20, random_state=42)
train, test = df.loc[train_ids].copy(), df.loc[test_ids].copy()
hp_median = float(train.horsepower.median())
cleaned = df.copy()
cleaned['horsepower'] = cleaned['horsepower'].fillna(hp_median)''')
p('Median imputation is transparent and less influenced by the high-power tail than a mean. It still concentrates several observations at one value, reduces uncertainty artificially and may weaken relationships with other variables. The missingness flag makes that treatment visible but does not prove why values were missing.')
h('Vehicle name standardization')
p('The program trims surrounding spaces, lowercases names and collapses repeated whitespace. It applies an explicit first-token alias map and handles the mercedes benz prefix before token replacement to avoid duplicating benz. The original car_name_raw and a row-level correction file are retained. These are conservative analyst-defined spelling and alias corrections, not a verified vehicle registry.')
table(['Raw example', 'Standardized value'], [
('chevy c20','chevrolet c20'),('toyouta corona mark ii (sw)','toyota corona mark ii (sw)'),
('maxda rx3','mazda rx3'),('vokswagen rabbit','volkswagen rabbit'),
('mercedes benz 300d','mercedes-benz 300d')],[3.4,3.4])
h('Erroneous entries and duplicates')
p('Conversion checks distinguish declared missing markers from unexpected nonnumeric tokens. Positive-value checks apply to MPG, displacement, horsepower, weight and acceleration. Allowed source codes are checked for cylinders, origin and model year. No violations or exact duplicates were detected. The program stops for review if these checks fail on a changed source rather than silently inventing corrections or dropping records.')
p('Repeated vehicle names alone are not duplicates: vehicles can share a name while differing in year or specifications. No rows were removed merely because their name appeared more than once.')

page('4 Outlier identification and treatment')
p('An outlier flag indicates an unusual value, not a confirmed error. For displacement, horsepower, weight and acceleration, the program learns quartiles from observed training values before imputation. It uses IQR = Q3 − Q1, with fences Q1 − 1.5 × IQR and Q3 + 1.5 × IQR. Missing horsepower is handled separately and is not treated as an outlier.')
code('''q1, q3 = train[col].quantile([.25, .75])
lo = q1 - 1.5 * (q3 - q1)
hi = q3 + 1.5 * (q3 - q1)
flagged = (df[col] < lo) | (df[col] > hi)''')
table(['Feature','Lower','Upper','Train','Test'],[[r['feature'],f"{r['lower']:.2f}",f"{r['upper']:.2f}",r['train_flagged'],r['test_flagged']] for r in a['outliers']],[1.8,1.2,1.2,1.3,1.3])
p('Horsepower produces 14 flags and acceleration produces 7. One record overlaps, giving 20 distinct flagged vehicles. Displacement and weight produce no flags. Negative statistical lower fences do not imply that negative physical measurements are permitted; separate domain checks enforce positivity.')
doc.add_picture(str(OUT/'outliers.png'),width=Inches(6.8))
p('Figure 2. Training-set boxplots show observations beyond the conventional 1.5 IQR whiskers. These values remain in the cleaned data.')
h('Decision and rationale')
p('All flagged values are retained because the available evidence does not establish that they are recording errors. Automatic removal would narrow the vehicle range and could bias later analysis toward typical vehicles. Capping would also change actual observations. The audit files expose each flag so later modelling can compare robust scaling or robust regression using training-only validation.')
p('MPG is not filtered using target-based fences, and discrete category codes are not screened with IQR rules. The saved outlier flags are review aids and are excluded from the 14 prepared predictors.')

page('5 Feature preparation without data leakage')
p('Data leakage occurs when information from the held-out test data influences fitting. The numerical imputer, scaler and category encoder are fitted only on the 318 training rows; transform is then applied to the 80 test rows. This follows the scikit-learn guidance on separate fitting and transformation. [3]')
h('Numerical and categorical processing')
p('Displacement, horsepower, weight, acceleration and model year use median imputation followed by standardization. Only horsepower actually requires a fill in this dataset. Standardization subtracts each training mean and divides by its training standard deviation. This puts predictors with different numerical scales on comparable footing.')
p('Cylinders and origin use one-hot encoding. Treating origin codes as continuous would impose an artificial numerical distance. Encoding cylinders as categories also avoids assuming that each extra cylinder has the same linear effect. All training categories are retained: five cylinder categories and three origin categories. Unknown future categories are encoded as zeros within their feature block using handle_unknown="ignore". [4]')
code('''transformer = ColumnTransformer([
    ('numeric', Pipeline([
        ('imputer', SimpleImputer(strategy='median')),
        ('scaler', StandardScaler())]), NUM),
    ('categorical', OneHotEncoder(
        handle_unknown='ignore', sparse_output=False), CAT),
    ('missing', 'passthrough', ['horsepower_missing'])])
x_train = transformer.fit_transform(train)
x_test = transformer.transform(test)
joblib.dump(transformer, out / 'preprocessor.joblib')''')
h('Feature accounting')
table(['Component','Feature count'],[('Scaled numerical predictors','5'),('Cylinder indicators','5'),('Origin indicators','3'),('Horsepower missing indicator','1'),('Total predictor columns','14')],[5.2,1.6])
p('The row_id column accompanies each CSV for alignment and must be dropped before training a model. Targets are supplied separately as mpg. Vehicle names, raw names, split labels and outlier review flags are excluded from the feature matrix. The fitted joblib file is a preprocessor, not a prediction model.')
p('Keeping every category can create redundant columns for an unregularized linear model with an intercept. A later model-specific pipeline can drop a reference category. The Week 1 representation retains the complete category information.')

page('6 Validation and analytical consequences')
h('Checks performed on the completed run')
table(['Validation','Outcome'],[
('Original records preserved','398 input and 398 cleaned rows'),
('Missing values in the nine cleaned source fields','0'),
('Training and test row identifiers overlap','0'),
('Prepared arrays contain only finite numbers','Passed'),
('Training numerical means approximately zero','Passed with tolerance 1e−10'),
('Reloaded preprocessor reproduces test features','Passed'),
('Feature columns match across partitions','14 predictors in both partitions')],[5.2,1.6])
code('''assert set(train.row_id).isdisjoint(test.row_id)
assert len(train) + len(test) == len(df)
assert np.isfinite(x_train).all() and np.isfinite(x_test).all()
assert cleaned[COLS].isna().sum().sum() == 0
assert np.allclose(
    joblib.load(out / 'preprocessor.joblib').transform(test),
    x_test)''')
h('Challenges and resolutions')
p('Quoted names and irregular separators could have shifted fields during parsing; a whitespace parser with quotation support and a strict shape check resolved that risk. Question marks could have left horsepower as text; explicit missing-marker replacement and numerical conversion produced a usable numeric column. Ambiguous text correction was limited to an auditable alias map, with originals preserved.')
p('Outlier rules identified unusual vehicles but could not establish an error. Keeping flagged rows preserved information while documenting uncertainty. Splitting before statistical preprocessing prevented the test set from influencing learned values. Bundling the source copy allows the normal run to work without a live download.')
h('Impact on later analysis')
p('Imputation preserves sample size but may understate variation and distort relationships. Standardization helps scale-sensitive methods; it does not make skewed data normally distributed and does not remove outliers. Retained extremes may still influence least-squares estimates. One-hot encoding avoids artificial category ordering but increases dimensionality and gives rare categories limited supporting data.')
p('The random split supports a demonstration within this historical dataset. It is not a test of performance on modern vehicles or future years. A chronological or vehicle-family split may be more appropriate for those questions. For cross-validation, refit the entire preprocessor inside each training fold using data before imputation; do not reuse these already prepared matrices across folds. Model selection and accuracy evaluation remain future work.')

page('7 Reproduction and references')
h('Run the submission')
p('Use Python 3.12 and extract the ZIP before running the commands from the project folder. The tested package versions are pinned in requirements.txt. The default command uses the included source file and regenerates the numerical outputs and figures. The separate report builder reads those outputs to recreate this Word document.')
code('''python -m venv .venv
# Windows activation
.venv\\Scripts\\activate
# macOS or Linux activation: source .venv/bin/activate
python -m pip install -r requirements.txt
python preprocess.py
python build_report.py
# Optional source refresh requiring internet
python preprocess.py --download''')
h('Essential files')
table(['File or folder','Purpose'],[
('Week1_Report.docx','Submission report with code excerpts and generated figures'),
('preprocess.py and build_report.py','Executable workflow and report regeneration'),
('requirements.txt and README.md','Dependencies, run instructions and output guidance'),
('data/raw and data/auto_mpg_raw.csv','UCI source files and parsed raw CSV'),
('outputs/auto_mpg_cleaned.csv','Readable cleaned data with provenance and split labels'),
('outputs/X_*.csv and y_*.csv','Prepared features and aligned targets'),
('outputs/preprocessor.joblib','Fitted preprocessing object'),
('outputs/audit_summary.json and other logs','Counts, checks, correction records and figures')],[3.1,3.7])
h('Source integrity')
p('The SHA256 checksum of the raw file used for this report is recorded below and in audit_summary.json. A changed checksum after a fresh download requires reviewing the updated source and regenerating the report.')
code(a['sha256'])
h('References')
for text in [
'[1] Quinlan, R. (1993). Auto MPG. UCI Machine Learning Repository. DOI: 10.24432/C5859H. https://archive.ics.uci.edu/dataset/9/auto+mpg',
'[2] Kaggle. Auto-mpg dataset. Alternative dataset listing; not the downloaded source. https://www.kaggle.com/datasets/uciml/autompg-dataset',
'[3] scikit-learn. Common pitfalls and recommended practices. https://scikit-learn.org/stable/common_pitfalls.html',
'[4] scikit-learn. OneHotEncoder API documentation. https://scikit-learn.org/stable/modules/generated/sklearn.preprocessing.OneHotEncoder.html']:
    para=doc.add_paragraph(text)
    for run in para.runs: run.font.size=Pt(8.5)
p('Web references accessed on 10 September 2026. Numerical findings are calculated from the supplied source file by preprocess.py.')
doc.core_properties.title = 'Auto MPG Data Acquisition Cleaning and Preprocessing'
doc.core_properties.subject = 'Week 1 data preparation report'
doc.core_properties.author = ''
for element in [doc.styles.element, doc.element]:
    for border in element.xpath('.//w:pBdr'):
        border.getparent().remove(border)
doc.save(ROOT / 'Week1_Report.docx')
print('Created Week1_Report.docx')
