"""
Build the Week 4 report: Supervised Learning — Breast Cancer Classification.
Reads results_summary.json, classification_report.txt, summary_statistics.csv
and the figures/ directory produced by breast_cancer_classification.py.
"""
from pathlib import Path
import json
import pandas as pd
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parent
FIG = ROOT / "figures"
results = json.loads((ROOT / "results_summary.json").read_text())
cls_report = (ROOT / "classification_report.txt").read_text()
stats = pd.read_csv(ROOT / "summary_statistics.csv", index_col=0)

# ── Document setup ────────────────────────────────────────────────────
doc = Document()
sec = doc.sections[0]
sec.page_width, sec.page_height = Inches(8.5), Inches(11)
sec.top_margin = sec.bottom_margin = Inches(0.7)
sec.left_margin = sec.right_margin = Inches(0.8)

for name in ["Normal", "Title", "Subtitle", "Heading 1", "Heading 2"]:
    style = doc.styles[name]
    style.font.name = "Calibri"
    style.font.color.rgb = RGBColor(0, 0, 0)
doc.styles["Normal"].font.size = Pt(10.5)
doc.styles["Normal"].paragraph_format.space_after = Pt(7)
doc.styles["Normal"].paragraph_format.line_spacing = 1.08
doc.styles["Title"].font.size = Pt(26)
doc.styles["Heading 1"].font.size = Pt(18)
doc.styles["Heading 2"].font.size = Pt(12)

header = sec.header.paragraphs[0]
header.text = "WEEK 4    |    SUPERVISED LEARNING — BREAST CANCER CLASSIFICATION"
header.runs[0].font.size = Pt(8)

footer = sec.footer.paragraphs[0]
footer.alignment = 2
footer.add_run("Week 4 report  |  ")
field = OxmlElement("w:fldSimple")
field.set(qn("w:instr"), "PAGE")
footer._p.append(field)


# ── Helpers ───────────────────────────────────────────────────────────
def p(text):
    doc.add_paragraph(text)

def h(text):
    doc.add_heading(text, 2)

def page(title):
    doc.add_page_break()
    doc.add_heading(title, 1)

def code(text):
    para = doc.add_paragraph()
    para.paragraph_format.space_after = Pt(9)
    para.paragraph_format.line_spacing = 1
    for i, line in enumerate(text.strip().splitlines()):
        run = para.add_run(("\n" if i else "") + line)
        run.font.name = "DejaVu Sans Mono"
        run.font.size = Pt(8)

def table(headers, rows, widths=None):
    t = doc.add_table(rows=1, cols=len(headers))
    t.autofit = False
    if widths:
        for col, width in zip(t.columns, widths):
            col.width = Inches(width)
    for c, text in zip(t.rows[0].cells, headers):
        c.text = str(text)
    repeat = OxmlElement("w:tblHeader")
    t.rows[0]._tr.get_or_add_trPr().append(repeat)
    for row in rows:
        for c, text in zip(t.add_row().cells, row):
            c.text = str(text)
    borders = OxmlElement("w:tblBorders")
    for side in ["top", "left", "bottom", "right", "insideH", "insideV"]:
        el = OxmlElement("w:" + side)
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), "4")
        el.set(qn("w:color"), "D9D9D9")
        borders.append(el)
    t._tbl.tblPr.append(borders)
    for i, row_obj in enumerate(t.rows):
        for j, c in enumerate(row_obj.cells):
            if widths:
                c.width = Inches(widths[j])
            props = c._tc.get_or_add_tcPr()
            shade = OxmlElement("w:shd")
            shade.set(qn("w:fill"), "DCE6EC" if i == 0 else "FFFFFF")
            props.append(shade)
            margins = OxmlElement("w:tcMar")
            for side in ["top", "left", "bottom", "right"]:
                el = OxmlElement("w:" + side)
                el.set(qn("w:w"), "75")
                el.set(qn("w:type"), "dxa")
                margins.append(el)
            props.append(margins)
            for para in c.paragraphs:
                para.paragraph_format.space_after = Pt(3)
                para.paragraph_format.space_before = Pt(3)
                for run in para.runs:
                    run.font.size = Pt(9)
                    run.bold = i == 0


def fig(filename, caption, width=6.8):
    path = FIG / filename
    if path.exists():
        doc.add_picture(str(path), width=Inches(width))
        cp = doc.add_paragraph(caption)
        cp.runs[0].italic = True
        cp.runs[0].font.size = Pt(9)


# ── Derived values ────────────────────────────────────────────────────
tm = results["test_metrics"]
cv = results["cv_results"]
cc = results["class_counts"]

# ====================================================================
#  COVER PAGE
# ====================================================================
doc.add_heading("Breast Cancer Classification\nUsing Supervised Learning", 0)
p("Week 4 Project Report")
p("Prepared for academic submission")
p("Author: Mayank")

h("Project objective")
p(
    "This project builds, evaluates, and compares three supervised learning "
    "classifiers — Logistic Regression, Random Forest, and SVM (RBF kernel) — "
    "to predict whether a breast tumour is Malignant or Benign using digitised "
    "fine-needle aspirate (FNA) features from the Breast Cancer Wisconsin "
    "(Diagnostic) dataset, originally sourced from the UCI Machine Learning "
    "Repository and bundled with scikit-learn."
)

h("Main results at a glance")
table(
    ["Measure", "Value"],
    [
        ("Dataset size", f'{results["dataset_shape"][0]} samples × {results["dataset_shape"][1]} features'),
        ("Class balance", f'{cc["benign"]} benign / {cc["malignant"]} malignant'),
        ("Missing values", str(results["missing_values"])),
        ("Best model", results["best_model_name"]),
        ("Test accuracy", f'{tm["accuracy"]:.4f}'),
        ("Test F1 score", f'{tm["f1"]:.4f}'),
        ("Test ROC-AUC", f'{tm["roc_auc"]:.4f}'),
    ],
    [3.1, 3.7],
)

h("Scope")
p(
    "The work covers data loading, exploratory analysis, domain-informed "
    "feature engineering, pipeline-based model training with stratified "
    "cross-validation, hyperparameter tuning via grid search, and final "
    "evaluation on a held-out test set. All preprocessing is encapsulated "
    "inside scikit-learn Pipelines to prevent data leakage."
)

# ====================================================================
#  SECTION 1 — Dataset and Exploration
# ====================================================================
page("1  Dataset Description and Exploratory Analysis")

h("1.1  Data source")
p(
    "The Breast Cancer Wisconsin (Diagnostic) dataset contains "
    f'{results["dataset_shape"][0]} observations, each described by 30 real-valued '
    "features computed from a digitised image of a fine-needle aspirate of a "
    "breast mass. These features describe characteristics of the cell nuclei "
    "present in the image, including radius, texture, perimeter, area, "
    "smoothness, compactness, concavity, concave points, symmetry and fractal "
    "dimension — each reported as mean, standard error and worst (largest) "
    "value. The target variable 'diagnosis' has two classes: Malignant (0) and "
    "Benign (1). The dataset has no missing values."
)

h("1.2  Class distribution")
fig("class_distribution.png",
    "Figure 1. Class distribution showing the balance between malignant and "
    "benign diagnoses. The dataset is moderately imbalanced, with benign cases "
    "outnumbering malignant cases roughly 1.7:1.")

h("1.3  Feature statistics (excerpt)")
# Show a subset of the summary stats
excerpt_features = ["mean radius", "mean texture", "mean perimeter", "mean area",
                    "mean concavity", "mean concave points", "worst radius", "worst area"]
stat_rows = []
for feat in excerpt_features:
    if feat in stats.index:
        stat_rows.append([
            feat,
            f'{stats.loc[feat, "mean"]:.3f}',
            f'{stats.loc[feat, "std"]:.3f}',
            f'{stats.loc[feat, "min"]:.3f}',
            f'{stats.loc[feat, "max"]:.3f}',
        ])
table(["Feature", "Mean", "Std Dev", "Min", "Max"], stat_rows, [2.2, 1.1, 1.1, 1.1, 1.1])
p("Only a representative subset is shown; the full summary is in summary_statistics.csv.")

h("1.4  Correlation heatmap")
fig("correlation_heatmap.png",
    "Figure 2. Correlation heatmap of the top features correlated with "
    "'mean radius'. High collinearity is expected among size-related features "
    "(radius, perimeter, area), motivating feature selection.",
    width=6.0)

h("1.5  Feature distributions by class")
fig("feature_boxplots.png",
    "Figure 3. Boxplots of mean radius and mean texture by diagnosis. "
    "Malignant tumours tend to have larger radius and higher texture values, "
    "confirming their discriminative power.")

# ====================================================================
#  SECTION 2 — Feature Engineering and Preprocessing
# ====================================================================
page("2  Feature Engineering and Preprocessing")

h("2.1  Engineered features")
p(
    "Three domain-informed ratio features were created to capture "
    "relationships between existing predictors:"
)
table(
    ["Feature", "Formula", "Rationale"],
    [
        ("radius_texture_ratio", "mean_radius / mean_texture",
         "Captures whether size outpaces surface irregularity"),
        ("concavity_area_ratio", "mean_concavity / mean_area",
         "Normalises concavity by overall mass size"),
        ("perimeter_radius_ratio", "mean_perimeter / mean_radius",
         "Detects shape deviations from a perfect circle"),
    ],
    [2.0, 2.3, 2.5],
)

code("""X[\"radius_texture_ratio\"] = X[\"mean radius\"] / (X[\"mean texture\"] + 1e-6)
X[\"concavity_area_ratio\"]  = X[\"mean concavity\"] / (X[\"mean area\"] + 1e-6)
X[\"perimeter_radius_ratio\"] = X[\"mean perimeter\"] / (X[\"mean radius\"] + 1e-6)""")

h("2.2  Train–test split")
p(
    "The data were split 80/20 using stratified sampling (random_state = 42) "
    "to preserve the class ratio in both partitions. The resulting training "
    "set contains 455 samples and the test set contains 114 samples."
)

h("2.3  Pipeline architecture")
p(
    "Each model is wrapped in a scikit-learn Pipeline comprising three stages: "
    "(1) StandardScaler for feature normalisation, "
    "(2) SelectKBest with ANOVA F-value for univariate feature selection "
    "(initial k = 15), and "
    "(3) the classifier itself. This design ensures all preprocessing is "
    "fitted only on training folds during cross-validation, preventing "
    "data leakage."
)
code("""Pipeline([
    (\"scaler\", StandardScaler()),
    (\"select\", SelectKBest(score_func=f_classif, k=15)),
    (\"clf\", LogisticRegression(max_iter=5000, random_state=42))
])""")

# ====================================================================
#  SECTION 3 — Model Training and Comparison
# ====================================================================
page("3  Model Training and Cross-Validation")

h("3.1  Models evaluated")
p(
    "Three classifiers were evaluated using 5-fold stratified "
    "cross-validation with F1 score as the primary metric (appropriate "
    "for the moderate class imbalance):"
)

cv_rows = []
for model_name, cv_data in cv.items():
    cv_rows.append([
        model_name,
        f'{cv_data["mean_f1"]:.4f}',
        f'±{cv_data["std_f1"]:.4f}',
    ])
table(["Model", "Mean CV F1", "Std Dev"], cv_rows, [2.5, 2.0, 2.0])

fig("model_comparison.png",
    "Figure 4. Cross-validated F1 scores across the three models. "
    "All three perform comparably, with Logistic Regression slightly ahead.")

h("3.2  Best model selection")
p(
    f'Based on cross-validation, {results["best_model_name"]} was selected '
    f'as the best candidate with a mean CV F1 of '
    f'{cv[results["best_model_name"]]["mean_f1"]:.4f}.'
)

# ====================================================================
#  SECTION 4 — Hyperparameter Tuning
# ====================================================================
page("4  Hyperparameter Tuning")

h("4.1  Grid search configuration")
p(
    f'GridSearchCV was applied to the {results["best_model_name"]} pipeline '
    "with the following parameter grid:"
)
table(
    ["Parameter", "Values tested"],
    [
        ("select__k (number of features)", "10, 15, 20"),
        ("clf__C (regularisation strength)", "0.01, 0.1, 1, 10"),
    ],
    [3.4, 3.4],
)

h("4.2  Best parameters")
bp = results["best_params"]
table(
    ["Parameter", "Optimal value"],
    [(k, str(v)) for k, v in bp.items()],
    [3.4, 3.4],
)
p(
    f'After tuning, the best cross-validated F1 score improved to '
    f'{results["best_cv_f1"]:.4f}. The model selected k = {bp.get("select__k", "N/A")} '
    f'features with regularisation C = {bp.get("clf__C", "N/A")}.'
)

# ====================================================================
#  SECTION 5 — Test Set Evaluation
# ====================================================================
page("5  Final Evaluation on Held-Out Test Set")

h("5.1  Performance metrics")
table(
    ["Metric", "Score"],
    [
        ("Accuracy", f'{tm["accuracy"]:.4f}'),
        ("Precision", f'{tm["precision"]:.4f}'),
        ("Recall", f'{tm["recall"]:.4f}'),
        ("F1 Score", f'{tm["f1"]:.4f}'),
        ("ROC-AUC", f'{tm["roc_auc"]:.4f}'),
    ],
    [3.4, 3.4],
)

h("5.2  Classification report")
code(cls_report)

h("5.3  Confusion matrix")
fig("confusion_matrix.png",
    f'Figure 5. Confusion matrix for {results["best_model_name"]} on the '
    "test set. The model achieves strong performance on both classes, "
    "with very few misclassifications.",
    width=4.5)

h("5.4  ROC curve")
fig("roc_curve.png",
    f'Figure 6. ROC curve for {results["best_model_name"]}. The area under '
    f'the curve (AUC = {tm["roc_auc"]:.3f}) indicates excellent '
    "discriminative ability, with the curve hugging the top-left corner.",
    width=4.5)

# ====================================================================
#  SECTION 6 — Feature Analysis
# ====================================================================
page("6  Feature Importance Analysis")

h("6.1  Selected features")
sel = results["selected_features"]
p(
    f"The final model selected {len(sel)} features via ANOVA F-value ranking. "
    "The selected features span size measurements (radius, perimeter, area), "
    "shape descriptors (compactness, concavity, concave points), and the "
    "engineered perimeter_radius_ratio, confirming that both raw and "
    "engineered features contribute to the model."
)

# list them in a compact table
feat_rows = [(f, "✓") for f in sel]
table(["Feature", "Selected"], feat_rows, [4.5, 1.5])

h("6.2  Feature importance / coefficient plot")
fig("feature_importance.png",
    f'Figure 7. Feature importances or coefficients from '
    f'{results["best_model_name"]}. Features related to worst-case '
    "measurements (worst radius, worst area, worst concave points) tend "
    "to carry the highest weight.",
    width=5.5)

# ====================================================================
#  SECTION 7 — Conclusions
# ====================================================================
page("7  Conclusions and Future Work")

h("7.1  Summary of findings")
p(
    f'The {results["best_model_name"]} classifier achieves a test accuracy of '
    f'{tm["accuracy"]:.2%} and an ROC-AUC of {tm["roc_auc"]:.4f} on the '
    "Breast Cancer Wisconsin dataset. The pipeline-based approach with "
    "stratified cross-validation and grid search ensured robust model "
    "selection without data leakage. Domain-informed feature engineering "
    "(ratio features) and ANOVA-based feature selection both contributed "
    "to the final model's performance."
)

h("7.2  Key takeaways")
p("1. All three classifiers (Logistic Regression, Random Forest, SVM) "
  "performed well on this dataset, indicating that the features are "
  "highly discriminative for the binary classification task.")
p("2. Encapsulating preprocessing inside Pipelines prevented data leakage "
  "and made the workflow reproducible.")
p("3. The engineered perimeter_radius_ratio feature was selected by the "
  "final model, validating the domain-informed feature engineering step.")

h("7.3  Limitations and future work")
p("• The dataset is relatively small (569 samples); performance on larger "
  "clinical datasets should be validated.")
p("• More advanced feature selection (e.g. recursive feature elimination) "
  "or dimensionality reduction (PCA) could be explored.")
p("• Ensemble methods such as XGBoost or stacking could potentially "
  "improve performance further.")
p("• The clinical applicability of this model requires external validation "
  "on independent hospital datasets.")

# ====================================================================
#  SECTION 8 — Reproduction
# ====================================================================
page("8  Reproduction and References")

h("How to run")
code("""python -m venv .venv
# Windows activation
.venv\\Scripts\\activate
pip install -r requirements.txt
python breast_cancer_classification.py
python build_report.py""")

h("Essential files")
table(
    ["File", "Purpose"],
    [
        ("Week4_Report.docx", "This submission report"),
        ("breast_cancer_classification.py", "Main pipeline script"),
        ("build_report.py", "Regenerates this report from saved outputs"),
        ("results_summary.json", "Serialised metrics and parameters"),
        ("classification_report.txt", "Scikit-learn classification report"),
        ("summary_statistics.csv", "Descriptive statistics of all features"),
        ("figures/", "All generated plots embedded in this report"),
    ],
    [3.2, 3.6],
)

h("References")
for text in [
    "[1] Street, W.N., Wolberg, W.H. and Mangasarian, O.L. (1993). "
    "Nuclear feature extraction for breast tumor diagnosis. "
    "IS&T/SPIE 1993 International Symposium on Electronic Imaging.",
    "[2] UCI Machine Learning Repository — Breast Cancer Wisconsin "
    "(Diagnostic). https://archive.ics.uci.edu/ml/datasets/Breast+Cancer+Wisconsin+(Diagnostic)",
    "[3] scikit-learn documentation: Pipelines and composite estimators. "
    "https://scikit-learn.org/stable/modules/compose.html",
]:
    para = doc.add_paragraph(text)
    for run in para.runs:
        run.font.size = Pt(8.5)

# ── Save ──────────────────────────────────────────────────────────────
doc.core_properties.title = "Breast Cancer Classification Using Supervised Learning"
doc.core_properties.subject = "Week 4 supervised learning report"
doc.core_properties.author = "Mayank"

# Clean stray borders
for element in [doc.styles.element, doc.element]:
    for border in element.xpath(".//w:pBdr"):
        border.getparent().remove(border)

doc.save(ROOT / "Week4_Report.docx")
print("Created Week4_Report.docx")
