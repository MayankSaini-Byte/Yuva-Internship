"""
Build the Week 6 report: Integrative Capstone — Diamond Pricing Analytics.
Reads results_summary.json, cluster_profile.csv, summary_statistics.csv
and the figures/ directory produced by diamond_capstone.py.
"""
from pathlib import Path
import json
import pandas as pd
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parent
FIG = ROOT / "figures"
results = json.loads((ROOT / "results_summary.json").read_text())
cluster_profile = pd.read_csv(ROOT / "cluster_profile.csv")
stats = pd.read_csv(ROOT / "summary_statistics.csv", index_col=0)

reg = results["regression"]
clust = results["clustering"]
cleaning = results["cleaning_log"]
tm = reg["test_metrics"]
cv = reg["cv_results"]

# ── Document setup ────────────────────────────────────────────────────
doc = Document()
sec = doc.sections[0]
sec.page_width, sec.page_height = Inches(8.5), Inches(11)
sec.top_margin = sec.bottom_margin = Inches(0.7)
sec.left_margin = sec.right_margin = Inches(0.8)

for name in ["Normal", "Title", "Subtitle", "Heading 1", "Heading 2"]:
    style = doc.styles[name]
    style.font.name = "Calibri"
    style.font.color.rgb = RGBColor(0, 0, 0)
doc.styles["Normal"].font.size = Pt(10.5)
doc.styles["Normal"].paragraph_format.space_after = Pt(7)
doc.styles["Normal"].paragraph_format.line_spacing = 1.08
doc.styles["Title"].font.size = Pt(26)
doc.styles["Heading 1"].font.size = Pt(18)
doc.styles["Heading 2"].font.size = Pt(12)

header = sec.header.paragraphs[0]
header.text = "WEEK 6    |    INTEGRATIVE CAPSTONE — DIAMOND PRICING ANALYTICS"
header.runs[0].font.size = Pt(8)

footer = sec.footer.paragraphs[0]
footer.alignment = 2
footer.add_run("Week 6 report  |  ")
field = OxmlElement("w:fldSimple")
field.set(qn("w:instr"), "PAGE")
footer._p.append(field)


# ── Helpers ───────────────────────────────────────────────────────────
def p(text):
    doc.add_paragraph(text)

def h(text):
    doc.add_heading(text, 2)

def page(title):
    doc.add_page_break()
    doc.add_heading(title, 1)

def code(text):
    para = doc.add_paragraph()
    para.paragraph_format.space_after = Pt(9)
    para.paragraph_format.line_spacing = 1
    for i, line in enumerate(text.strip().splitlines()):
        run = para.add_run(("\n" if i else "") + line)
        run.font.name = "DejaVu Sans Mono"
        run.font.size = Pt(8)

def table(headers, rows, widths=None):
    t = doc.add_table(rows=1, cols=len(headers))
    t.autofit = False
    if widths:
        for col, width in zip(t.columns, widths):
            col.width = Inches(width)
    for c, text in zip(t.rows[0].cells, headers):
        c.text = str(text)
    repeat = OxmlElement("w:tblHeader")
    t.rows[0]._tr.get_or_add_trPr().append(repeat)
    for row in rows:
        for c, text in zip(t.add_row().cells, row):
            c.text = str(text)
    borders = OxmlElement("w:tblBorders")
    for side in ["top", "left", "bottom", "right", "insideH", "insideV"]:
        el = OxmlElement("w:" + side)
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), "4")
        el.set(qn("w:color"), "D9D9D9")
        borders.append(el)
    t._tbl.tblPr.append(borders)
    for i, row_obj in enumerate(t.rows):
        for j, c in enumerate(row_obj.cells):
            if widths:
                c.width = Inches(widths[j])
            props = c._tc.get_or_add_tcPr()
            shade = OxmlElement("w:shd")
            shade.set(qn("w:fill"), "DCE6EC" if i == 0 else "FFFFFF")
            props.append(shade)
            margins = OxmlElement("w:tcMar")
            for side in ["top", "left", "bottom", "right"]:
                el = OxmlElement("w:" + side)
                el.set(qn("w:w"), "75")
                el.set(qn("w:type"), "dxa")
                margins.append(el)
            props.append(margins)
            for para in c.paragraphs:
                para.paragraph_format.space_after = Pt(3)
                para.paragraph_format.space_before = Pt(3)
                for run in para.runs:
                    run.font.size = Pt(9)
                    run.bold = i == 0
    doc.add_paragraph().paragraph_format.space_after = Pt(1)


def fig(filename, caption, width=6.8):
    path = FIG / filename
    if path.exists():
        doc.add_picture(str(path), width=Inches(width))
        cp = doc.add_paragraph(caption)
        cp.runs[0].italic = True
        cp.runs[0].font.size = Pt(9)


# ====================================================================
#  COVER PAGE
# ====================================================================
doc.add_heading(
    "Diamond Pricing Analytics\nPredictive Pricing and Market Segmentation", 0
)
p("Week 6 — Integrative Capstone Project Report")
p("Prepared for academic submission")
p("Author: Mayank")

h("Project objective")
p(
    "This capstone project integrates the full data-science pipeline — from "
    "data acquisition and cleaning through exploratory analysis, feature "
    "engineering, supervised regression modelling, and unsupervised "
    "clustering — to build a diamond price prediction system and discover "
    "natural market segments within the diamond trade. The project uses the "
    "publicly available 'diamonds' dataset (~54 000 records) from the "
    "seaborn-data repository."
)

h("Main results at a glance")
table(
    ["Measure", "Value"],
    [
        ("Raw dataset size", f'{cleaning["rows_before"]:,} records'),
        ("Cleaned dataset size", f'{cleaning["rows_after"]:,} records'),
        ("Best regression model", reg["best_model"]),
        ("Test R²", f'{tm["r2"]:.4f}'),
        ("Test RMSE", f'${tm["rmse"]:.2f}'),
        ("Test MAE", f'${tm["mae"]:.2f}'),
        ("Market segments discovered", str(clust["chosen_k"])),
        ("Clustering silhouette score", f'{clust["final_silhouette"]:.4f}'),
    ],
    [3.1, 3.7],
)

h("Scope")
p(
    "This report documents all eight pipeline stages: data acquisition, "
    "data cleaning, exploratory data analysis, feature engineering, "
    "supervised price prediction (regression), unsupervised market "
    "segmentation (K-Means clustering), evaluation, and insight generation. "
    "Both supervised and unsupervised techniques are combined in a single "
    "integrative workflow."
)

# ====================================================================
#  SECTION 1 — Data Acquisition and Cleaning
# ====================================================================
page("1  Data Acquisition and Cleaning")

h("1.1  Data source")
p(
    "The diamonds dataset contains records of round-cut diamond "
    "transactions, with each record described by 10 attributes: carat "
    "weight, cut quality (Fair to Ideal), color grade (D–J), clarity "
    "grade (I1–IF), depth percentage, table width, price (USD), and "
    "three physical dimension measurements (x, y, z in mm). The dataset "
    "is widely used in data science education and benchmarking."
)

h("1.2  Cleaning steps")
table(
    ["Cleaning step", "Records affected"],
    [
        ("Exact duplicate removal", str(cleaning["duplicates_removed"])),
        ("Invalid zero-dimension rows", str(cleaning["invalid_zero_dimension_rows_removed"])),
        ("Physical outliers (domain bounds)", str(cleaning["physical_outliers_removed"])),
        ("Remaining missing values", str(cleaning["missing_values_remaining"])),
    ],
    [3.8, 3.0],
)
p(
    f'After cleaning, {cleaning["rows_after"]:,} records remain from the '
    f'original {cleaning["rows_before"]:,}. Duplicate removal, zero-dimension '
    "filtering, and domain-informed outlier bounds (e.g. y > 20 mm, z > 10 mm, "
    "table outside 40–80, depth outside 40–80) were applied sequentially."
)
code("""df = df.drop_duplicates()
df = df[(df[\"x\"] > 0) & (df[\"y\"] > 0) & (df[\"z\"] > 0)]
outlier_mask = (df[\"y\"] > 20) | (df[\"z\"] > 10) | \\
               (df[\"table\"] > 80) | (df[\"table\"] < 40) | \\
               (df[\"depth\"] > 80) | (df[\"depth\"] < 40)
df = df[~outlier_mask]""")

# ====================================================================
#  SECTION 2 — Exploratory Data Analysis
# ====================================================================
page("2  Exploratory Data Analysis")

h("2.1  Price distribution")
fig("price_distribution.png",
    "Figure 1. Price distribution in raw scale (left) and log-transformed "
    "scale (right). The raw distribution is right-skewed, with the bulk "
    "of diamonds priced under $5,000 and a long tail extending to $18,000+. "
    "Log transformation reveals a bimodal pattern.")

h("2.2  Price vs. carat")
fig("price_vs_carat.png",
    "Figure 2. Scatter plot of price vs. carat (4,000-record sample, "
    "colored by cut quality). The strong nonlinear relationship between "
    "carat and price is evident, with higher-quality cuts clustering at "
    "various price points.")

h("2.3  Categorical price breakdown")
fig("categorical_price_breakdown.png",
    "Figure 3. Mean price by cut, color, and clarity grades. "
    "Counter-intuitively, 'Fair' cut diamonds sometimes show higher mean "
    "prices because larger stones tend to receive more conservative cuts. "
    "This confound motivates controlling for carat in the regression model.",
    width=6.5)

h("2.4  Correlation heatmap")
fig("correlation_heatmap.png",
    "Figure 4. Correlation matrix of numeric features. Carat, x, y, z, "
    "and price are highly intercorrelated (all r > 0.85), confirming that "
    "physical size dominates price. Depth and table show weaker correlations.",
    width=5.5)

# ====================================================================
#  SECTION 3 — Feature Engineering
# ====================================================================
page("3  Feature Engineering")

h("3.1  Engineered features")
p("Five features were created or encoded to improve model performance:")
table(
    ["Feature", "Type", "Description"],
    [
        ("volume", "Engineered", "x × y × z — physical volume proxy"),
        ("price_per_carat", "Engineered", "price / carat — unit value (used for clustering)"),
        ("cut_enc", "Ordinal", "Fair=0 → Ideal=4"),
        ("color_enc", "Ordinal", "J=0 (worst) → D=6 (best)"),
        ("clarity_enc", "Ordinal", "I1=0 (worst) → IF=7 (best)"),
    ],
    [1.7, 1.3, 3.8],
)
code("""df_fe[\"volume\"] = df_fe[\"x\"] * df_fe[\"y\"] * df_fe[\"z\"]
df_fe[\"price_per_carat\"] = df_fe[\"price\"] / df_fe[\"carat\"]

cut_order = [\"Fair\", \"Good\", \"Very Good\", \"Premium\", \"Ideal\"]
color_order = [\"J\", \"I\", \"H\", \"G\", \"F\", \"E\", \"D\"]
clarity_order = [\"I1\", \"SI2\", \"SI1\", \"VS2\", \"VS1\", \"VVS2\", \"VVS1\", \"IF\"]""")

h("3.2  Final feature set for regression")
p(
    f'The regression model uses {len(reg["feature_columns"])} features: '
    + ", ".join(reg["feature_columns"])
    + "."
)

# ====================================================================
#  SECTION 4 — Supervised Modelling (Regression)
# ====================================================================
page("4  Supervised Modelling — Price Prediction")

h("4.1  Models evaluated")
p(
    "Four regression models were compared using 3-fold cross-validation "
    "(on a 6,000-sample subsample for computational efficiency) with R² "
    "as the evaluation metric. All models include a StandardScaler step."
)
cv_rows = []
for model_name, cv_data in cv.items():
    cv_rows.append([
        model_name,
        f'{cv_data["mean_r2"]:.4f}',
        f'±{cv_data["std_r2"]:.4f}',
    ])
table(["Model", "Mean CV R²", "Std Dev"], cv_rows, [2.5, 2.0, 2.0])

fig("regression_model_comparison.png",
    "Figure 5. Cross-validated R² scores for all four regression models. "
    "Tree-based models (Random Forest, Gradient Boosting) clearly outperform "
    "linear models, capturing the nonlinear price–feature relationships.",
    width=5.5)

h("4.2  Model selection and tuning")
p(
    f'{reg["best_model"]} was selected as the best model with a CV R² of '
    f'{reg["best_cv_r2"]:.4f}. A lightweight grid search was performed:'
)
bp = reg["best_params"]
if bp:
    table(
        ["Parameter", "Optimal value"],
        [(k, str(v)) for k, v in bp.items()],
        [3.4, 3.4],
    )

h("4.3  Test set evaluation")
table(
    ["Metric", "Score"],
    [
        ("R²", f'{tm["r2"]:.4f}'),
        ("RMSE", f'${tm["rmse"]:.2f}'),
        ("MAE", f'${tm["mae"]:.2f}'),
    ],
    [3.4, 3.4],
)
p(
    f'The {reg["best_model"]} explains {tm["r2"]:.1%} of the price variance '
    f'on held-out test data, with an average prediction error of '
    f'${tm["mae"]:.0f} (MAE).'
)

h("4.4  Actual vs. predicted prices")
fig("actual_vs_predicted.png",
    f'Figure 6. Actual vs. predicted prices for {reg["best_model"]}. '
    "Points cluster tightly around the diagonal, with some dispersion "
    "at higher price points where fewer training samples exist.",
    width=5.0)

h("4.5  Residual analysis")
fig("residual_plot.png",
    "Figure 7. Residual plot showing prediction errors vs. predicted price. "
    "The increasing spread at higher prices indicates mild heteroscedasticity, "
    "suggesting that predicting premium diamonds is inherently harder.",
    width=5.0)

h("4.6  Feature importance")
fig("feature_importance.png",
    f'Figure 8. Feature importances from {reg["best_model"]}. '
    "Carat (and its correlated volume/dimension features) dominates, "
    "followed by clarity and color encodings.",
    width=5.0)

# ====================================================================
#  SECTION 5 — Unsupervised Modelling (Clustering)
# ====================================================================
page("5  Unsupervised Modelling — Market Segmentation")

h("5.1  Clustering approach")
p(
    "K-Means clustering was applied to an 8,000-sample subset using "
    "seven features: carat, depth, table, cut_enc, color_enc, clarity_enc, "
    "and price_per_carat. All features were standardised before clustering."
)

h("5.2  Optimal cluster selection")
fig("cluster_selection.png",
    "Figure 9. Elbow method (left) and silhouette analysis (right) for "
    "selecting the number of clusters. While k = 2 maximises the silhouette "
    "score, a 2-segment view has little practical value for merchandising.",
    width=6.5)

p(
    f'The silhouette-optimal k is {clust["silhouette_optimal_k"]}, but a '
    f'business-informed decision selected k = {clust["chosen_k"]} to provide '
    "actionable market segments. The margin between the two was within 0.02, "
    "making the trade-off acceptable."
)

h("5.3  Cluster profiles")
# Build cluster profile table from the CSV
profile_headers = list(cluster_profile.columns)
profile_rows = []
for _, row in cluster_profile.iterrows():
    profile_rows.append([str(row[col]) for col in profile_headers])
table(profile_headers, profile_rows)

fig("cluster_profile_heatmap.png",
    "Figure 10. Cluster profile heatmap showing z-score normalised feature "
    "levels (color intensity) with actual mean values annotated. This "
    "reveals distinct segment characteristics.",
    width=6.0)

h("5.4  PCA visualization")
fig("cluster_pca.png",
    f'Figure 11. PCA projection of the {clust["chosen_k"]} clusters. '
    "The first two principal components capture the dominant variation "
    "(primarily driven by carat/size). Clusters show reasonable separation "
    "in this reduced space.",
    width=5.5)

h("5.5  Segment interpretation")
p("Based on the cluster profiles, the four market segments can be characterised as:")
segment_labels = {
    0: ("Mid-range standard", "Medium carat (~0.67), moderate quality grades, "
        "average price-per-carat. The everyday consumer segment."),
    1: ("Budget small stones", "Small carat (~0.46), high cut quality (Ideal-dominant), "
        "good clarity. Entry-level, value-oriented buyers."),
    2: ("Premium large stones", "Large carat (~1.37), moderate quality, highest absolute "
        "price ($8,981 avg). Luxury/investment segment."),
    3: ("Deep-cut mid-range", "Medium carat (~0.80), deepest depth (63.6%), lower cut "
        "quality. Conservative cuts on mid-size stones."),
}
seg_rows = [(str(k), label, desc) for k, (label, desc) in segment_labels.items()]
table(["Cluster", "Label", "Description"], seg_rows, [0.8, 2.0, 4.0])

# ====================================================================
#  SECTION 6 — Integration and Insights
# ====================================================================
page("6  Integration and Insights")

h("6.1  How supervised and unsupervised results complement each other")
p(
    "The regression model quantifies how each attribute contributes to price, "
    "while the clustering reveals that the market naturally stratifies into "
    "segments with distinct feature profiles. Together, they enable "
    "segment-specific pricing strategies: for example, the premium large-stone "
    "segment (Cluster 2) shows higher residual variance, suggesting that "
    "within that segment, additional factors (brand, certification, "
    "fluorescence) likely play a larger role."
)

h("6.2  Business recommendations")
p("1. Pricing engine: Deploy the Gradient Boosting model as a baseline price "
  "estimator for incoming inventory. The $346 MAE provides a useful "
  "benchmark for negotiation and listing.")
p("2. Inventory segmentation: Use the 4-cluster segmentation to tailor "
  "marketing strategies — budget buyers (Cluster 1) respond to value "
  "messaging, while premium buyers (Cluster 2) value rarity and size.")
p("3. Quality vs. size trade-offs: The counter-intuitive EDA finding — "
  "that lower-cut grades sometimes show higher mean prices — should be "
  "communicated to sales teams with the carat-confound explanation.")

# ====================================================================
#  SECTION 7 — Conclusions
# ====================================================================
page("7  Conclusions and Future Work")

h("7.1  Summary")
p(
    f'This capstone project delivered a {reg["best_model"]} regression model '
    f'achieving R² = {tm["r2"]:.4f} on held-out test data, and a K-Means '
    f'clustering model that identified {clust["chosen_k"]} actionable market '
    f'segments with a silhouette score of {clust["final_silhouette"]:.4f}. '
    "The end-to-end pipeline covers all stages from raw data to insights."
)

h("7.2  Limitations")
p("• The model does not account for diamond certification body, "
  "fluorescence, or cut proportions beyond depth and table.")
p("• Clustering on a subsample may miss rare segments. Full-dataset "
  "clustering or mini-batch K-Means could address this.")
p("• Heteroscedasticity in residuals suggests that a log-price model "
  "or quantile regression might improve predictions for high-value stones.")

h("7.3  Future work")
p("• Incorporate external data (GIA reports, auction records) for richer "
  "feature sets.")
p("• Experiment with XGBoost, LightGBM, or neural network regressors.")
p("• Apply DBSCAN or Gaussian Mixture Models for non-spherical cluster "
  "discovery.")
p("• Build an interactive dashboard (Streamlit/Dash) for real-time "
  "pricing and segment assignment.")

# ====================================================================
#  SECTION 8 — Reproduction and References
# ====================================================================
page("8  Reproduction and References")

h("How to run")
code("""python -m venv .venv
# Windows activation
.venv\\Scripts\\activate
pip install -r requirements.txt
python diamond_capstone.py
python build_report.py""")

h("Essential files")
table(
    ["File", "Purpose"],
    [
        ("Week6_Report.docx", "This submission report"),
        ("diamond_capstone.py", "Complete capstone pipeline script"),
        ("build_report.py", "Regenerates this report from saved outputs"),
        ("results_summary.json", "All metrics, parameters, and clustering results"),
        ("cluster_profile.csv", "Mean feature values per cluster"),
        ("summary_statistics.csv", "Descriptive statistics"),
        ("figures/", "All generated plots embedded in this report"),
    ],
    [3.2, 3.6],
)

h("References")
for text in [
    "[1] Wickham, H. (2016). ggplot2: Elegant Graphics for Data Analysis. "
    "Springer-Verlag New York. (Original diamonds dataset source.)",
    "[2] scikit-learn documentation: KMeans clustering. "
    "https://scikit-learn.org/stable/modules/clustering.html#k-means",
    "[3] scikit-learn documentation: Gradient Boosting Regressor. "
    "https://scikit-learn.org/stable/modules/ensemble.html#gradient-boosted-regression-trees",
    "[4] Rousseeuw, P.J. (1987). Silhouettes: a graphical aid to the "
    "interpretation and validation of cluster analysis. Journal of "
    "Computational and Applied Mathematics, 20, 53–65.",
]:
    para = doc.add_paragraph(text)
    for run in para.runs:
        run.font.size = Pt(8.5)

# ── Save ──────────────────────────────────────────────────────────────
doc.core_properties.title = "Diamond Pricing Analytics — Predictive Pricing and Market Segmentation"
doc.core_properties.subject = "Week 6 integrative capstone report"
doc.core_properties.author = "Mayank"

# Clean stray borders
for element in [doc.styles.element, doc.element]:
    for border in element.xpath(".//w:pBdr"):
        border.getparent().remove(border)

doc.save(ROOT / "Week6_Report.docx")
print("Created Week6_Report.docx")
